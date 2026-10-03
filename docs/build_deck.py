"""
BAHI — Round 1 deck.

Ten slides: the seven sections the organisers mandate, in their order, then
three value-adds (the four rules, the demo and its risks, the sources).

Three rules drive it.

1. SKIMMED, NOT READ. 1,771 teams registered and the PDF alone decides the
   shortlist, so a reviewer gives each page seconds. Every slide makes one
   point, its title is that point as a sentence, and the body stays near 100
   words at 12pt or larger.

2. ONE PICTURE PER SLIDE. A flow, an architecture, phone screens or a table,
   drawn as native shapes so the PDF stays sharp and the file stays editable.

3. EVERY FACT IS SOURCED, AND ONLY USEFUL FACTS ARE KEPT. A number, a claim
   about another company or a statement of law carries its source on the same
   slide. Our own design is stated as design, our demo data is labelled as
   demo data, and anything we could not source was cut.

It should look like Paytm: pale sky ground, floating white cards, navy
headings, cyan as a rule or mark and never as text (#00BAF2 is ~2.0:1 on
white). Arial, because LibreOffice's PDF export silently swaps .ttc fonts
like Helvetica Neue for a serif.
"""

from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches as I, Pt

# ── Paytm palette ──────────────────────────────────────────────────────────
NAVY = RGBColor(0x00, 0x29, 0x70)
CYAN = RGBColor(0x00, 0xBA, 0xF2)
CYAN_TX = RGBColor(0x00, 0x77, 0xA8)
SKY = RGBColor(0xEA, 0xF6, 0xFC)
SKY_2 = RGBColor(0xD6, 0xEE, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x1C, 0x1C, 0x1C)
MUTED = RGBColor(0x5B, 0x66, 0x75)
FAINT = RGBColor(0x8B, 0x95, 0xA5)
HAIR = RGBColor(0xD9, 0xE4, 0xEC)
GREEN = RGBColor(0x1E, 0x7E, 0x34)
RED = RGBColor(0xB3, 0x26, 0x1E)
PALE = RGBColor(0x8F, 0xA4, 0xC6)

F = "Arial"
LEFT, CENTER = PP_ALIGN.LEFT, PP_ALIGN.CENTER
MID = MSO_ANCHOR.MIDDLE

W, H = I(13.333), I(7.5)
ML = I(0.62)
CW = W - I(1.24)
TOP = I(1.80)          # where content starts under the title
BOTTOM = I(6.72)       # where content must end, above the footer

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]
TOTAL = 10


# ── primitives ─────────────────────────────────────────────────────────────
def flat(shape):
    """Drop the theme style python-pptx attaches, which LibreOffice renders as
    a drop shadow on every shape (and thickens every hairline)."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)
    return shape


def slide(bg=SKY):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, W, H, bg)
    return s


def rect(s, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, radius=None):
    r = s.shapes.add_shape(shape, x, y, w, h)
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    flat(r)
    if radius is not None:
        r.adjustments[0] = radius
    return r


def card(s, x, y, w, h, color=WHITE, radius=0.06):
    return rect(s, x, y, w, h, color, MSO_SHAPE.ROUNDED_RECTANGLE, radius)


def tb(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tf = s.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def para(tf, text, size, color=INK, bold=False, line=1.3, after=0, before=0,
         first=False, align=None, spc=None, link=None):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.line_spacing = line
    p.space_after = Pt(after)
    p.space_before = Pt(before)
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    if link:
        r.hyperlink.address = link
    f = r.font
    f.size, f.name, f.bold = Pt(size), F, bold
    f.color.rgb = color
    if spc is not None:
        f._rPr.set("spc", str(int(spc * 100)))
    return p


def rich(tf, parts, size, first=False, line=1.3, after=0, before=0, align=None):
    """One paragraph of mixed runs: parts are (text, color, bold)."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.line_spacing = line
    p.space_after = Pt(after)
    p.space_before = Pt(before)
    if align is not None:
        p.alignment = align
    for text, color, bold in parts:
        r = p.add_run()
        r.text = text
        r.font.size, r.font.name, r.font.bold = Pt(size), F, bold
        r.font.color.rgb = color
    return p


def eyebrow(s, text):
    para(tb(s, ML, I(0.42), CW, I(0.25)), text.upper(), 10, CYAN_TX,
         bold=True, first=True, spc=1.4)


def title(s, text, size=26):
    para(tb(s, ML, I(0.74), CW, I(0.9)), text, size, NAVY, bold=True,
         first=True, line=1.05)


def footer(s, text=None):
    """Source line bottom-left, page number bottom-right."""
    n = len(prs.slides._sldIdLst)
    if text:
        para(tb(s, ML, I(6.98), CW - I(0.8), I(0.3)), text, 9, FAINT,
             first=True, line=1.2)
    para(tb(s, W - ML - I(0.7), I(6.98), I(0.7), I(0.3)), f"{n} / {TOTAL}",
         9, FAINT, first=True, align=PP_ALIGN.RIGHT)


def arrow(s, x1, y1, x2, y2, color=NAVY, width=2.0, dashed=False):
    c = flat(s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    if dashed:
        c.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    end = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
    end.set("type", "triangle")
    end.set("w", "med")
    end.set("len", "med")
    return c


def num(s, x, y, n, d=I(0.36), fill=NAVY, color=WHITE, size=12):
    rect(s, x, y, d, d, fill, MSO_SHAPE.OVAL)
    para(tb(s, x, y, d, d, anchor=MID), str(n), size, color, bold=True,
         first=True, align=CENTER, line=1.0)


def pill(s, x, y, text, fill, color, w, h=I(0.26), size=9):
    rect(s, x, y, w, h, fill, MSO_SHAPE.ROUNDED_RECTANGLE, 0.5)
    para(tb(s, x, y, w, h, anchor=MID), text, size, color, bold=True,
         first=True, align=CENTER, line=1.0, spc=0.8)


def node(s, x, y, w, h, head, sub=None, fill=WHITE, head_color=NAVY,
         sub_color=MUTED, head_size=13, sub_size=11, outline=None, dashed=False):
    r = card(s, x, y, w, h, fill, 0.12)
    if outline is not None:
        r.line.color.rgb = outline
        r.line.width = Pt(1.25)
        if dashed:
            r.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tf = tb(s, x + I(0.12), y, w - I(0.24), h, anchor=MID)
    para(tf, head, head_size, head_color, bold=True, first=True, align=CENTER,
         line=1.1)
    if sub:
        para(tf, sub, sub_size, sub_color, align=CENTER, line=1.2, before=3)


def phone(s, path, x, y, w):
    """The renders are 1312×2592; keep that ratio."""
    s.shapes.add_picture(path, x, y, width=w, height=int(w * 2592 / 1312))


def cols(n, gutter, x=ML, width=CW):
    w = int((width - gutter * (n - 1)) / n)
    return [(x + i * (w + gutter), w) for i in range(n)]


# ═══════════════════════════════════════════════════════════════════════════
# 1 · TITLE
# ═══════════════════════════════════════════════════════════════════════════
s = slide(NAVY)
rect(s, 0, 0, W, Pt(8), CYAN)

tf = tb(s, ML, I(1.55), I(7.0), I(4.6))
para(tf, "BAHI", 70, WHITE, bold=True, first=True, line=0.95)
para(tf, "Udhaar, confirmed by both sides.", 28, WHITE, bold=True, before=10,
     line=1.1)
para(tf, "The shopkeeper speaks a credit entry to his Paytm Soundbox. The "
         "customer confirms it on his own phone and pays in one tap. Every "
         "night, an agent decides who actually needs a reminder.",
     15, RGBColor(0xC9, 0xD8, 0xEE), before=18, line=1.4)
rect(s, ML, I(5.55), I(0.5), Pt(3), CYAN)
tf = tb(s, ML, I(5.75), I(7.3), I(0.8))
para(tf, "Paytm Build for India AI Hackathon  ·  Mumbai, 3 October 2026",
     12, WHITE, first=True)
para(tf, "Track 2: AI-Powered Financial Journeys", 12, WHITE, before=4)
para(tf, "Team Bitzlab", 12, PALE, before=10)

phone(s, "docs/screens/A2-speak.png", I(8.25), I(1.25), I(2.25))
phone(s, "docs/screens/B1-confirm.png", I(10.72), I(1.25), I(2.25))
para(tb(s, I(8.25), I(5.86), I(4.72), I(0.3)),
     "Shopkeeper’s screen  ·  Customer’s screen  ·  demo data", 10, PALE,
     first=True, align=CENTER)

# ═══════════════════════════════════════════════════════════════════════════
# 2 · PROBLEM
# ═══════════════════════════════════════════════════════════════════════════
s = slide()
eyebrow(s, "1 · Problem statement")
title(s, "The udhaar book has one author. The customer never sees it.")

card(s, ML, TOP, CW, I(2.05))
ny, nh, nw = I(2.20), I(1.25), I(2.7)
xs = [ML + I(0.45), ML + I(4.69), ML + I(8.94)]
node(s, xs[0], ny, nw, nh, "Shopkeeper", "writes every entry", fill=SKY,
     head_size=15, sub_size=12)
node(s, xs[1], ny, nw, nh, "The udhaar book", "a notebook, or a ledger app",
     fill=SKY, head_size=15, sub_size=12)
node(s, xs[2], ny, nw, nh, "Customer", "no copy  ·  no total  ·  no proof",
     fill=WHITE, head_color=MUTED, outline=FAINT, dashed=True, head_size=15,
     sub_size=12)
ay = ny + nh // 2
arrow(s, xs[0] + nw + I(0.1), ay, xs[1] - I(0.1), ay)
arrow(s, xs[1] + nw + I(0.1), ay, xs[2] - I(0.1), ay, FAINT, dashed=True)
para(tb(s, xs[0] + nw, ay - I(0.36), xs[1] - xs[0] - nw, I(0.3)), "writes",
     11, NAVY, bold=True, first=True, align=CENTER)
para(tb(s, xs[1] + nw, ay - I(0.36), xs[2] - xs[1] - nw, I(0.3)),
     "an SMS, at best", 11, MUTED, first=True, align=CENTER)

points = [
    ("He can’t see patterns",
     "Someone who always pays on day 9 and someone quietly drifting away look "
     "the same in a notebook. So everyone is chased the same way."),
    ("The customer has no record",
     "He can’t see what he owes across shops. If the book says ₹500 and he "
     "remembers ₹300, he has nothing to show."),
    ("Apps went digital on one side only",
     "Khatabook alone has 5 Cr+ downloads. Its customers still get only an SMS "
     "— no app of their own, and no say in the entry."),
]
for (x, w), (head, body) in zip(cols(3, I(0.4)), points):
    rect(s, x, I(4.28), I(0.4), Pt(3), CYAN)
    tf = tb(s, x, I(4.45), w, I(2.2))
    para(tf, head, 16, NAVY, bold=True, first=True, line=1.15)
    para(tf, body, 13, MUTED, before=8, line=1.4)
footer(s, "Source: Khatabook on Google Play (5 Cr+ downloads) and khatabook.com, "
          "checked 23 Sep 2026.")

# ═══════════════════════════════════════════════════════════════════════════
# 3 · SOLUTION
# ═══════════════════════════════════════════════════════════════════════════
s = slide()
eyebrow(s, "2 · Proposed solution")
title(s, "The same book, now confirmed by both sides.")

steps = [
    ("docs/screens/A2-speak.png", "He speaks it",
     "“Sharma ko do sau udhaar” is shown back to check before it is sent."),
    ("docs/screens/B1-confirm.png", "The customer confirms",
     "He taps “Yes, I owe ₹200” in his own Paytm app."),
    ("docs/screens/B2-mybook.png", "He sees every shop",
     "One total across all his shops, and one tap to pay."),
    ("docs/screens/A3-tonight.png", "The agent picks the day",
     "Sharma always pays on day 9, so on day 4 he hears nothing."),
]
PW = I(1.95)
for i, ((x, w), (path, head, body)) in enumerate(zip(cols(4, I(0.28)), steps)):
    phone(s, path, x + (w - PW) // 2, I(1.66), PW)
    num(s, x + I(0.05), I(5.64), i + 1)
    tf = tb(s, x + I(0.52), I(5.66), w - I(0.52), I(1.2))
    para(tf, head, 14, NAVY, bold=True, first=True, line=1.1)
    para(tf, body, 12, MUTED, before=4, line=1.3)
footer(s, "Screens from our working prototype. Names and amounts are demo data.")

# ═══════════════════════════════════════════════════════════════════════════
# 4 · TECH
# ═══════════════════════════════════════════════════════════════════════════
s = slide()
eyebrow(s, "3 · Technology / tech stack")
title(s, "AI hears the words. Plain code decides every number.")


def tagged(s, x, y, w, h, head, sub, tag=None):
    node(s, x, y, w, h, head, sub, head_size=13, sub_size=11)
    if tag == "AI":
        pill(s, x + I(0.14), y - I(0.13), "AI", CYAN, NAVY, I(0.46))
    elif tag == "CODE":
        pill(s, x + I(0.14), y - I(0.13), "CODE", NAVY, WHITE, I(0.66))


def row(s, y, label, items):
    para(tb(s, ML, y - I(0.40), CW, I(0.25)), label.upper(), 9.5, MUTED,
         bold=True, first=True, spc=1.2)
    n = len(items)
    gap = I(0.52)
    w = int((CW - gap * (n - 1)) / n)
    h = I(1.05)
    for i, (head, sub, tag) in enumerate(items):
        x = ML + i * (w + gap)
        tagged(s, x, y, w, h, head, sub, tag)
        if i < n - 1:
            arrow(s, x + w + I(0.07), y + h // 2, x + w + gap - I(0.07),
                  y + h // 2, width=1.75)


row(s, I(2.12), "During the day — recording an entry", [
    ("Shopkeeper speaks", "to the Soundbox", None),
    ("Sarvam", "speech → text", "AI"),
    ("“Sharma ko do sau udhaar”", "text only", None),
    ("Rule-based parser", "text → ₹200", "CODE"),
    ("Ledger", "Postgres 17", "CODE"),
])
row(s, I(3.95), "Every night — deciding who hears from us", [
    ("n8n schedule", "runs at 11 pm", None),
    ("Timing engine", "each customer’s own repayment gap, in SQL", "CODE"),
    ("Cognee memory", "the shopkeeper’s notes: “Patil pays after the 10th”",
     "AI"),
    ("A reminder — or silence", "most nights, silence", None),
])

card(s, ML, I(5.35), CW, I(0.78))
tf = tb(s, ML + I(0.25), I(5.35), CW - I(0.5), I(0.78), anchor=MID)
rich(tf, [("How we prove it:  ", NAVY, True),
          ("a test fails the build if the core logic imports any AI or "
           "database code. The model only ever returns text.", INK, False)],
     13, first=True)

stack = ["Next.js 16", "React 19", "FastAPI", "Python 3.13", "Postgres 17",
         "Sarvam", "Cognee", "n8n"]
x = ML
para(tb(s, x, I(6.36), I(0.8), I(0.3), anchor=MID), "STACK", 9.5, MUTED,
     bold=True, first=True, spc=1.2)
x += I(0.8)
for name in stack:
    w = I(0.28 + 0.085 * len(name))
    pill(s, x, I(6.36), name, WHITE, NAVY, w, I(0.3), size=10.5)
    x += w + I(0.12)
footer(s)

# ═══════════════════════════════════════════════════════════════════════════
# 5 · USP
# ═══════════════════════════════════════════════════════════════════════════
s = slide()
eyebrow(s, "4 · Unique selling proposition")
title(s, "Every udhaar app serves the shopkeeper. We add the customer.")

COLW = [I(3.0), I(2.6), I(3.35), I(3.14)]
HEADS = ["", "Downloads", "Customer has an app", "Customer confirms entries"]
ROWS = [
    ("Khatabook", "5 Cr+", "No — SMS / WhatsApp", "No"),
    ("OkCredit", "1 Cr+", "No — “They don’t need the app”", "No"),
    ("Paytm Business Khata", "Inside Paytm for Business",
     "No — SMS with a pay link", "No"),
    ("Done", "100+", "Yes", "Yes — by OTP"),
    ("BAHI", "—", "Yes — inside the Paytm app", "Yes — one tap"),
]
y = TOP
x = ML
for w, head in zip(COLW, HEADS):
    para(tb(s, x, y, w, I(0.3)), head.upper(), 9.5, CYAN_TX, bold=True,
         first=True, spc=1.0)
    x += w
y += I(0.34)
rect(s, ML, y, CW, Pt(1.5), CYAN)
y += I(0.08)
RH = I(0.5)
for r, cells in enumerate(ROWS):
    ours = r == len(ROWS) - 1
    if ours:
        card(s, ML - I(0.1), y, CW + I(0.2), RH, WHITE, 0.2)
    x = ML
    for c, (w, val) in enumerate(zip(COLW, cells)):
        tf = tb(s, x + (I(0.02) if c else 0), y, w - I(0.1), RH, anchor=MID)
        if c == 0:
            para(tf, val, 14, NAVY if ours else INK, bold=True, first=True)
        elif c == 1:
            para(tf, val, 12.5, MUTED, first=True, bold=ours)
        else:
            color = GREEN if val.startswith("Yes") else RED
            para(tf, val, 12.5, color, first=True, bold=ours)
        x += w
    y += RH
    if not ours:
        rect(s, ML, y - Pt(0.5), CW, Pt(0.75), HAIR)

usps = [
    ("Two-sided", "An entry becomes a fact only when the customer confirms it."),
    ("Voice-first", "Spoken to the Soundbox already on the counter. No typing."),
    ("Timed per person",
     "Reminders follow each customer’s own repayment pattern."),
]
for (x, w), (head, body) in zip(cols(3, I(0.3)), usps):
    card(s, x, I(5.05), w, I(1.55))
    rect(s, x + I(0.25), I(5.28), I(0.4), Pt(3), CYAN)
    tf = tb(s, x + I(0.25), I(5.42), w - I(0.5), I(1.1))
    para(tf, head, 15, NAVY, bold=True, first=True)
    para(tf, body, 12.5, MUTED, before=5, line=1.35)
footer(s, "Sources: Google Play listings, okcredit.in, khatabook.com and "
          "business.paytm.com, checked 23 Sep 2026.")

# ═══════════════════════════════════════════════════════════════════════════
# 6 · IMPACT
# ═══════════════════════════════════════════════════════════════════════════
s = slide()
eyebrow(s, "5 · Impact & benefits")
title(s, "Fewer reminders, sent only to people who are actually late.")

LW = I(3.75)
card(s, ML, TOP, LW, I(4.85), NAVY)
tf = tb(s, ML + I(0.35), TOP + I(0.45), LW - I(0.7), I(4.2))
para(tf, "4 of 38", 54, WHITE, bold=True, first=True, line=1.0)
para(tf, "customers who owe get a reminder tonight.", 15, WHITE, bold=True,
     before=14, line=1.3)
para(tf, "The other 34 are not late by their own pattern, so they hear "
         "nothing.", 13, RGBColor(0xC9, 0xD8, 0xEE), before=12, line=1.4)
para(tb(s, ML + I(0.35), TOP + I(4.2), LW - I(0.7), I(0.5)),
     "Our demo shop: 60 customers, six months of synthetic data.", 10, PALE,
     first=True, line=1.3)

who = [
    ("For the shopkeeper", ["Stops chasing people who always pay.",
                            "Spots the ones who are drifting.",
                            "Records by voice, so the queue keeps moving."]),
    ("For the customer", ["Sees his total across every shop.",
                          "Has proof of every entry.",
                          "Is reminded privately, not at the counter."]),
    ("For Paytm", ["Udhaar settled in cash moves to UPI.",
                   "A daily reason to use the Soundbox.",
                   "A new reason to open the Paytm app."]),
]
for (x, w), (head, lines) in zip(cols(3, I(0.26), ML + LW + I(0.3),
                                      CW - LW - I(0.3)), who):
    card(s, x, TOP, w, I(4.85))
    rect(s, x + I(0.25), TOP + I(0.28), I(0.4), Pt(3), CYAN)
    tf = tb(s, x + I(0.25), TOP + I(0.45), w - I(0.5), I(4.2))
    para(tf, head, 15, NAVY, bold=True, first=True, after=6)
    for ln in lines:
        para(tf, ln, 13, INK, before=12, line=1.35)
footer(s, "These are the outcomes we expect; a pilot would measure them. "
          "“4 of 38” is computed from our synthetic demo shop.")

# ═══════════════════════════════════════════════════════════════════════════
# 7 · BUSINESS MODEL
# ═══════════════════════════════════════════════════════════════════════════
s = slide()
eyebrow(s, "6 · Business model")
title(s, "Free for both sides. Paytm gains through the Soundbox.")

LW = I(7.45)
card(s, ML, TOP, LW, I(2.3))
fy, fh, fw = TOP + I(0.85), I(1.0), I(1.6)
fx = [ML + I(0.3), ML + I(2.92), ML + I(5.55)]
node(s, fx[0], fy, fw, fh, "Customer", "pays nothing extra", fill=SKY,
     head_size=14, sub_size=11)
node(s, fx[1], fy, fw, fh, "Shopkeeper", "pays nothing extra", fill=SKY,
     head_size=14, sub_size=11)
node(s, fx[2], fy, fw, fh, "Paytm", "UPI + Soundbox", fill=NAVY,
     head_color=WHITE, sub_color=RGBColor(0xC9, 0xD8, 0xEE), head_size=14,
     sub_size=11)
for a, b, label in [(0, 1, "clears udhaar\nby UPI"),
                    (1, 2, "keeps his\nSoundbox")]:
    ay = fy + fh // 2
    arrow(s, fx[a] + fw + I(0.08), ay, fx[b] - I(0.08), ay)
    tf = tb(s, fx[a] + fw, fy - I(0.52), fx[b] - fx[a] - fw, I(0.5),
            anchor=MSO_ANCHOR.BOTTOM)
    for k, part in enumerate(label.split("\n")):
        para(tf, part, 10.5, NAVY, bold=True, first=(k == 0), align=CENTER,
             line=1.1)

tf = tb(s, ML, I(4.40), LW, I(2.3))
para(tf, "What Paytm gets", 15, NAVY, bold=True, first=True, after=2)
for ln in ["Udhaar repayments move from cash to Paytm UPI.",
           "A reason for merchants to keep their Soundbox subscription.",
           "Later, with consent: a repayment history for people with no credit "
           "file. We only record it — licensed partners would lend."]:
    rich(tf, [("—  ", CYAN_TX, True), (ln, INK, False)], 13, before=9,
         line=1.35)

RX = ML + LW + I(0.3)
RW = CW - LW - I(0.3)
facts = [
    ("1.57 Cr", "storefronts already run a Paytm Soundbox — the counter BAHI "
                "lives on.",
     "Paytm Q1 FY27 earnings release, Jul 2026"),
    ("₹116 Cr", "Khatabook’s FY24 loss, on ₹102.7 Cr revenue. Charging for a "
                "ledger has not worked, so we don’t.",
     "YourStory, Nov 2024"),
]
for k, (big, body, src) in enumerate(facts):
    y = TOP + k * I(2.55)
    card(s, RX, y, RW, I(2.3))
    tf = tb(s, RX + I(0.3), y + I(0.25), RW - I(0.6), I(1.9))
    para(tf, big, 36, NAVY, bold=True, first=True, line=1.0)
    para(tf, body, 13, INK, before=8, line=1.35)
    para(tf, "Source: " + src, 9.5, FAINT, before=8)
footer(s)

# ═══════════════════════════════════════════════════════════════════════════
# 8 · FOUR RULES
# ═══════════════════════════════════════════════════════════════════════════
s = slide()
eyebrow(s, "Designed around Indian law")
title(s, "Four things BAHI refuses to do, each from Indian law.")

rules = [
    ("No due dates",
     "Customers tap “Yes, I owe ₹200” — never “I’ll pay by the 10th”.",
     "A dated promise can be a promissory note, which the IT Act does not "
     "cover.",
     "IT Act 2000, s.1(4) & First Schedule"),
    ("No fees, no interest",
     "Never a rupee above the price. There is no fee field anywhere.",
     "Money-lending laws define a loan by the interest it carries.",
     "e.g. Karnataka Money Lenders Act 1961, s.2(9)"),
    ("No debt kept alive forever",
     "We never ask a customer to re-confirm an old debt just to keep it alive.",
     "Each written acknowledgment restarts the 3-year limit.",
     "Limitation Act 1963, s.18"),
    ("No shaming, no pressure",
     "No defaulter lists, no public alerts, no “demand” button.",
     "Public labels risk defamation; repeated nudges for gain are “nagging”.",
     "BNS 2023, s.356  ·  CCPA Dark Patterns Guidelines 2023"),
]
CH = BOTTOM - TOP
for i, ((x, w), (head, body, why, cite)) in enumerate(zip(cols(4, I(0.24)),
                                                           rules)):
    card(s, x, TOP, w, CH)
    num(s, x + I(0.25), TOP + I(0.28), i + 1, I(0.46), size=15)
    tf = tb(s, x + I(0.25), TOP + I(0.95), w - I(0.5), I(2.3))
    para(tf, head, 16, NAVY, bold=True, first=True, line=1.12)
    para(tf, body, 13, INK, before=10, line=1.38)
    rect(s, x + I(0.25), TOP + I(3.25), w - I(0.5), Pt(0.75), HAIR)
    tf = tb(s, x + I(0.25), TOP + I(3.4), w - I(0.5), I(1.45))
    rich(tf, [("Why:  ", CYAN_TX, True), (why, MUTED, False)], 12, first=True,
         line=1.35)
    para(tf, cite, 10, CYAN_TX, bold=True, before=8, line=1.25)
footer(s, "Each rule is enforced in the database schema and the API, not only "
          "on screen. Full citations on slide 10.")

# ═══════════════════════════════════════════════════════════════════════════
# 9 · DEMO AND RISKS
# ═══════════════════════════════════════════════════════════════════════════
s = slide()
eyebrow(s, "On 3 October")
title(s, "A 40-second demo, and what could go wrong.")

LW = I(4.75)
card(s, ML, TOP, LW, BOTTOM - TOP)
tf = tb(s, ML + I(0.3), TOP + I(0.28), LW - I(0.6), I(0.4))
para(tf, "The demo", 15, NAVY, bold=True, first=True)
demo = ["The shopkeeper speaks an entry.",
        "A judge’s phone buzzes.",
        "The judge taps Confirm.",
        "Both books update, live.",
        "The agent says who gets a reminder tonight, and why."]
y = TOP + I(0.85)
for i, ln in enumerate(demo):
    num(s, ML + I(0.3), y, i + 1, I(0.34), size=11)
    para(tb(s, ML + I(0.8), y - I(0.02), LW - I(1.1), I(0.62)), ln, 13, INK,
         first=True, line=1.3)
    y += I(0.6)
para(tb(s, ML + I(0.3), BOTTOM - I(0.72), LW - I(0.6), I(0.6)),
     "All data is synthetic and labelled on screen. A browser microphone "
     "stands in for the Soundbox.", 10.5, FAINT, first=True, line=1.3)

RX = ML + LW + I(0.3)
RW = CW - LW - I(0.3)
C1 = I(2.45)
para(tb(s, RX, TOP, C1, I(0.3)), "RISK", 9.5, CYAN_TX, bold=True, first=True,
     spc=1.0)
para(tb(s, RX + C1, TOP, RW - C1, I(0.3)), "HOW WE HANDLE IT", 9.5, CYAN_TX,
     bold=True, first=True, spc=1.0)
rect(s, RX, TOP + I(0.34), RW, Pt(1.5), CYAN)
risks = [
    ("Speech mishears the amount",
     "The amount is shown before sending; “Say it again” is one tap. No model "
     "ever writes the number."),
    ("Venue wifi fails",
     "Deployed publicly, so the judge’s phone uses its own data. Speech is "
     "cached offline."),
    ("The customer never confirms",
     "The entry stays in the shopkeeper’s book, marked unconfirmed."),
    ("Reminders feel like pressure",
     "Timed per person, private and rare. There is no demand button."),
]
y = TOP + I(0.5)
RH = I(1.08)
for k, (risk, fix) in enumerate(risks):
    para(tb(s, RX, y, C1 - I(0.2), RH), risk, 14, NAVY, bold=True, first=True,
         line=1.2)
    para(tb(s, RX + C1, y, RW - C1, RH), fix, 13, MUTED, first=True, line=1.35)
    y += RH
    if k < len(risks) - 1:
        rect(s, RX, y - I(0.12), RW, Pt(0.75), HAIR)
footer(s)

# ═══════════════════════════════════════════════════════════════════════════
# 10 · SOURCES
# ═══════════════════════════════════════════════════════════════════════════
s = slide()
eyebrow(s, "Sources")
title(s, "Where every fact in this deck comes from.")

market = [
    ("Khatabook — 5 Cr+ downloads; customers get SMS / WhatsApp",
     "Google Play, checked 23 Sep 2026",
     "https://play.google.com/store/apps/details?id=com.vaibhavkalpe.android.khatabook"),
    ("OkCredit — 1 Cr+ downloads; “They don’t need the app”",
     "Google Play and okcredit.in, checked 23 Sep 2026",
     "https://okcredit.in"),
    ("Done — 100+ downloads; entries confirmed by OTP",
     "Google Play, checked 23 Sep 2026",
     "https://play.google.com/store/apps/details?id=com.doneapp.cpma"),
    ("Paytm Business Khata — customer reminders by SMS with a pay link",
     "business.paytm.com, checked 23 Sep 2026",
     "https://business.paytm.com/business-khata"),
    ("Paytm Soundbox “deployed at 1.57 Cr storefronts”",
     "Paytm Q1 FY2027 earnings release, 20 Jul 2026",
     "https://paytm.com/document/ir/financial-results/fy2026-27/"
     "Paytm_Earning-Release_Q1-FY-2027_INR.pdf"),
    ("Khatabook FY24 — ₹102.7 Cr revenue, ₹116 Cr loss",
     "YourStory, Nov 2024",
     "https://yourstory.com/2024/11/ms-dhoni-backed-khatabook-clocks-rs-1027-"
     "cr-revenue-cuts-losses-7-in-fy24"),
]
law = [
    ("Information Technology Act 2000, s.1(4) and First Schedule, entry 1 "
     "(as amended 2022)", "Negotiable instruments other than cheques are "
     "outside the Act."),
    ("Karnataka Money Lenders Act 1961, s.2(9)",
     "A “loan” is an advance at interest, of money or in kind."),
    ("Limitation Act 1963, s.18",
     "A written acknowledgment starts a fresh limitation period."),
    ("Bharatiya Nyaya Sanhita 2023, s.356", "Defamation."),
    ("CCPA Guidelines for Prevention and Regulation of Dark Patterns 2023, "
     "Annexure 1, item 10", "“Nagging”."),
]
(lx, lw), (rx, rw) = cols(2, I(0.4))
card(s, lx, TOP, lw, BOTTOM - TOP)
card(s, rx, TOP, rw, BOTTOM - TOP)

tf = tb(s, lx + I(0.3), TOP + I(0.25), lw - I(0.6), BOTTOM - TOP - I(0.4))
para(tf, "MARKET AND PRODUCTS", 9.5, CYAN_TX, bold=True, first=True, spc=1.2,
     after=2)
for fact, where, url in market:
    para(tf, fact, 11.5, INK, bold=True, before=9, line=1.2)
    para(tf, where, 10, MUTED, line=1.2, before=1, link=url)

tf = tb(s, rx + I(0.3), TOP + I(0.25), rw - I(0.6), BOTTOM - TOP - I(0.4))
para(tf, "LAW", 9.5, CYAN_TX, bold=True, first=True, spc=1.2, after=2)
for cite, gist in law:
    para(tf, cite, 11.5, INK, bold=True, before=9, line=1.2)
    para(tf, gist, 10, MUTED, line=1.2, before=1)
footer(s, "Every shop and customer name in this deck is synthetic demo data. "
          "Click a source line to open it.")

# Links in the theme are pure blue and purple; make them Paytm's text cyan.
theme = prs.slide_masters[0].part.part_related_by(RT.THEME)
theme._blob = (theme.blob
               .replace(b'<a:hlink><a:srgbClr val="0000FF"/>',
                        b'<a:hlink><a:srgbClr val="0077A8"/>')
               .replace(b'<a:folHlink><a:srgbClr val="800080"/>',
                        b'<a:folHlink><a:srgbClr val="0077A8"/>'))

out = "docs/BAHI-Round1.pptx"
prs.save(out)
print("wrote", out, "·", len(prs.slides._sldIdLst), "slides")
